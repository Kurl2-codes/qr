import docx
import sys

def read_docx(file_path):
    doc = docx.Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return '\n'.join(full_text)

if __name__ == "__main__":
    path = r'c:\Users\Raven Butial\Documents\QR Code\Guardrail\Guardrail.docx'
    out_path = r'c:\Users\Raven Butial\Documents\QR Code\guardrail_utf8.txt'
    try:
        doc = docx.Document(path)
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(f"Number of paragraphs: {len(doc.paragraphs)}\n")
            for i, para in enumerate(doc.paragraphs):
                if para.text.strip():
                    f.write(f"P{i}: {para.text}\n")
        print(f"Successfully wrote to {out_path}")
    except Exception as e:
        print(f"Error: {e}")
